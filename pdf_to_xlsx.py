#!/usr/bin/env python3
"""PDF -> XLSX. Wyciaga tabele z PDF, takze te bez linii siatki
(dane ulozone "wierszami" z drobnymi odstepami). Obsluga naglowka i stopki.

Uruchomienie: python pdf_to_xlsx.py            (GUI)
              python pdf_to_xlsx.py --selftest (test logiki)
"""
import json
import os
import re
import sys
import threading
import traceback
from bisect import bisect_right
from collections import Counter

# ---------------------------------------------------------------- logika ----


def cluster_rows(words, ytol=3.0):
    """Grupuje slowa w wiersze po pionowym srodku znaku (tolerancja ytol pt)."""
    rows = []  # [srodek, [slowa]]
    for w in sorted(words, key=lambda w: (w["top"], w["x0"])):
        c = (w["top"] + w["bottom"]) / 2
        if rows and abs(c - rows[-1][0]) <= ytol:
            r = rows[-1]
            r[1].append(w)
            r[0] = sum((x["top"] + x["bottom"]) / 2 for x in r[1]) / len(r[1])
        else:
            rows.append([c, [w]])
    return [sorted(ws, key=lambda w: w["x0"]) for _, ws in rows]


def column_bounds(rows, page_width, min_gap=6.0, max_cover=0.15):
    """Granice kolumn = srodki pionowych przerw (>= min_gap pt) wolnych
    w prawie wszystkich wierszach. max_cover dopuszcza pojedyncze wiersze
    ciagnace sie przez cala szerokosc (naglowek, tytul, suma)."""
    width = int(page_width) + 2
    occ = [0] * width
    for row in rows:
        marked = bytearray(width)
        for w in row:
            for x in range(max(0, int(w["x0"])), min(int(w["x1"]) + 1, width)):
                marked[x] = 1
        for x, v in enumerate(marked):
            occ[x] += v
    thr = int(len(rows) * max_cover)
    bounds, run, seen = [], 0, False
    for x, v in enumerate(occ):
        if v > thr:
            if seen and run >= min_gap:
                bounds.append(x - run / 2)
            seen, run = True, 0
        else:
            run += 1
    return bounds


def rows_from_words(page, ytol, min_gap):
    words = page.extract_words()
    if not words:
        return []
    rows_w = cluster_rows(words, ytol)
    bounds = column_bounds(rows_w, page.width, min_gap)
    out = []
    for row in rows_w:
        cells = [""] * (len(bounds) + 1)
        for w in row:
            i = bisect_right(bounds, (w["x0"] + w["x1"]) / 2)
            cells[i] = (cells[i] + " " + w["text"]).strip()
        out.append(cells)
    return out


def page_rows(page, mode, ytol, min_gap):
    if mode != "wiersze":
        tables = page.extract_tables()
        rows = [[(c or "").strip() for c in r] for t in tables for r in t]
        if rows or mode == "tabele":
            return rows
    return rows_from_words(page, ytol, min_gap)


def drop_repeats(pages_rows, min_frac=0.6):
    """Usuwa naglowki/stopki: wiersze z gory lub z dolu strony, ktore na
    wiekszosci stron stoja na tej samej pozycji. Wiersze ze srodka strony
    (dane) zostaja, nawet jesli sie powtarzaja. Zwraca (strony, usuniete)."""
    n = len(pages_rows)
    if n < 3:
        return pages_rows, []
    seen = Counter()
    for rows in pages_rows:
        for i, r in enumerate(rows):
            seen["t", i, tuple(r)] += 1
            seen["b", len(rows) - 1 - i, tuple(r)] += 1
    thr = max(2, int(n * min_frac))
    edges = []
    for rows in pages_rows:
        a, b = 0, len(rows)
        while a < b and seen["t", a, tuple(rows[a])] >= thr:
            a += 1
        # stopka: na KAZDEJ stronie (ostatnia strona listy zwykle konczy sie
        # innym wierszem danych niz pelne strony)
        while b > a and seen["b", len(rows) - b, tuple(rows[b - 1])] >= n:
            b -= 1
        edges.append((a, b))
    # wiersz, ktory gdziekolwiek wystepuje tez w srodku strony, to dane
    middle = {tuple(r) for rows, (a, b) in zip(pages_rows, edges) for r in rows[a:b]}
    out, removed = [], []
    for rows, (a, b) in zip(pages_rows, edges):
        keep = []
        for i, r in enumerate(rows):
            (keep if a <= i < b or tuple(r) in middle else removed).append(r)
        out.append(keep)
    return out, removed


NUM = re.compile(r"^-?\d{1,3}(?:[  ]\d{3})*(?:[.,]\d+)?$|^-?\d+(?:[.,]\d+)?$")
ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def cast(v, numbers):
    s = ILLEGAL.sub("", v).strip()
    digits = re.sub(r"\D", "", re.split(r"[.,]", s)[0])
    # numery kont (>15 cyfr Excel traci) i kody z zerem wiodacym zostaja tekstem

    if len(digits) > 15 or (len(digits) > 1 and digits[0] == "0"):
        return s
    if numbers and NUM.match(s):
        try:
            return float(s.replace(" ", "").replace(" ", "").replace(",", "."))
        except ValueError:
            pass
    return s


def convert(pdf_path, out_dir, o, log=print):
    import pdfplumber
    from openpyxl import Workbook

    pages_rows = []
    with pdfplumber.open(pdf_path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            top, bot = o["top"], o["bottom"]
            if top or bot:
                page = page.crop((0, top, page.width, max(top + 1, page.height - bot)))
            rows = [r for r in page_rows(page, o["mode"], o["ytol"], o["min_gap"])
                    if any(c.strip() for c in r)]
            pages_rows.append(rows)
            log("  strona %d: %d wierszy" % (i, len(rows)))

    if o["drop_repeats"]:
        pages_rows, removed = drop_repeats(pages_rows)
        log("  usunieto powtarzalne naglowki/stopki: %d wierszy" % len(removed))
        for r in dict.fromkeys(map(tuple, removed)):
            log("    - " + " | ".join(r))

    wb = Workbook()
    wb.remove(wb.active)
    if o["one_sheet"]:
        sheets = [("Dane", [r for rows in pages_rows for r in rows])]
    else:
        sheets = [("Strona %d" % i, rows)
                  for i, rows in enumerate(pages_rows, 1) if rows]
    for name, rows in sheets:
        ws = wb.create_sheet(name[:31])
        width = max((len(r) for r in rows), default=1)
        for r in rows:
            ws.append([cast(c, o["numbers"]) for c in r] + [""] * (width - len(r)))
    if not wb.sheetnames:
        wb.create_sheet("Dane")

    out = os.path.join(out_dir,
                       os.path.splitext(os.path.basename(pdf_path))[0] + ".xlsx")
    wb.save(out)
    return out


def run_batch(inp, out_dir, o, log=print):
    if os.path.isfile(inp):
        files = [inp]
    else:
        files = sorted(os.path.join(inp, f) for f in os.listdir(inp)
                       if f.lower().endswith(".pdf"))
    if not files:
        log("Brak plikow PDF w: %s" % inp)
        return
    os.makedirs(out_dir, exist_ok=True)
    for f in files:
        log("Przetwarzam: %s" % os.path.basename(f))
        try:
            log("OK -> %s" % convert(f, out_dir, o, log))
        except Exception:
            log("BLAD: %s\n%s" % (os.path.basename(f), traceback.format_exc()))
    log("Zakonczono (%d plikow)." % len(files))


# -------------------------------------------------------------------- GUI ----


def gui():
    import tkinter as tk
    from tkinter import filedialog, ttk, scrolledtext

    root = tk.Tk()
    root.title("PDF -> XLSX")
    root.geometry("780x580")
    pad = dict(padx=6, pady=3)

    cfg = load_config()
    v_in = tk.StringVar(value=cfg.get("last_in", ""))
    v_out = tk.StringVar(value=cfg.get("last_out", ""))
    v_mode = tk.StringVar(value="auto")
    v_top = tk.StringVar(value="0")
    v_bot = tk.StringVar(value="0")
    v_ytol = tk.StringVar(value="3")
    v_gap = tk.StringVar(value="6")
    v_one = tk.BooleanVar(value=True)
    v_rep = tk.BooleanVar(value=True)
    v_num = tk.BooleanVar(value=True)

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)

    ttk.Label(f, text="Plik PDF lub folder:").grid(row=0, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_in, width=60).grid(row=0, column=1, columnspan=3, **pad)
    ttk.Button(f, text="Plik...", command=lambda: v_in.set(
        filedialog.askopenfilename(filetypes=[("PDF", "*.pdf")]) or v_in.get())
    ).grid(row=0, column=4, **pad)
    ttk.Button(f, text="Folder...", command=lambda: v_in.set(
        filedialog.askdirectory() or v_in.get())).grid(row=0, column=5, **pad)

    ttk.Label(f, text="Folder wyjsciowy:").grid(row=1, column=0, sticky="w", **pad)
    ttk.Entry(f, textvariable=v_out, width=60).grid(row=1, column=1, columnspan=3, **pad)
    ttk.Button(f, text="Wybierz...", command=lambda: v_out.set(
        filedialog.askdirectory() or v_out.get())).grid(row=1, column=4, **pad)

    ttk.Label(f, text="Tryb:").grid(row=2, column=0, sticky="w", **pad)
    ttk.Combobox(f, textvariable=v_mode, width=10, state="readonly",
                 values=["auto", "tabele", "wiersze"]).grid(row=2, column=1,
                                                            sticky="w", **pad)
    ttk.Label(f, text="auto: linie tabeli, a gdy ich brak - podzial po wspolrzednych"
              ).grid(row=2, column=2, columnspan=4, sticky="w", **pad)

    g = ttk.LabelFrame(root, text="Naglowek / stopka i podzial")
    g.pack(fill="x", **pad)
    opts = [("Obetnij gore [pt]:", v_top, "wysokosc naglowka"),
            ("Obetnij dol [pt]:", v_bot, "wysokosc stopki"),
            ("Tolerancja wiersza [pt]:", v_ytol, "scalanie linii w jeden wiersz"),
            ("Min. przerwa kolumny [pt]:", v_gap, "podzial na kolumny")]
    for i, (lab, var, hint) in enumerate(opts):
        r, c = i // 2, (i % 2) * 3
        ttk.Label(g, text=lab).grid(row=r, column=c, sticky="w", **pad)
        ttk.Entry(g, textvariable=var, width=7).grid(row=r, column=c + 1, **pad)
        ttk.Label(g, text=hint, foreground="#666").grid(row=r, column=c + 2,
                                                        sticky="w", **pad)

    c = ttk.Frame(root)
    c.pack(fill="x", **pad)
    ttk.Checkbutton(c, text="Wszystkie strony na jednym arkuszu",
                    variable=v_one).pack(side="left", **pad)
    ttk.Checkbutton(c, text="Usun powtarzalne naglowki/stopki",
                    variable=v_rep).pack(side="left", **pad)
    ttk.Checkbutton(c, text="Konwertuj liczby", variable=v_num).pack(side="left", **pad)

    log_box = scrolledtext.ScrolledText(root, height=18)
    log_box.pack(fill="both", expand=True, **pad)

    def log(msg):
        def put():
            log_box.insert("end", str(msg) + "\n")
            log_box.see("end")
        root.after(0, put)

    btn = ttk.Button(root, text="Konwertuj")
    btn.pack(pady=6)

    def start():
        inp = v_in.get().strip('" ')
        out = v_out.get().strip('" ')
        if not inp or not os.path.exists(inp):
            return log("Wskaz istniejacy plik PDF lub folder.")
        if not out:
            return log("Wskaz folder wyjsciowy.")

        def num(var, d):
            try:
                return float(var.get().replace(",", "."))
            except ValueError:
                return d

        o = dict(mode=v_mode.get(), top=num(v_top, 0), bottom=num(v_bot, 0),
                 ytol=num(v_ytol, 3), min_gap=num(v_gap, 6),
                 one_sheet=v_one.get(), drop_repeats=v_rep.get(),
                 numbers=v_num.get())
        btn.config(state="disabled")
        log_box.delete("1.0", "end")

        def work():
            try:
                run_batch(inp, out, o, log)
            finally:
                root.after(0, lambda: btn.config(state="normal"))

        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)

    def on_close():
        save_config({"last_in": v_in.get(), "last_out": v_out.get()})
        root.destroy()

    root.protocol("WM_DELETE_WINDOW", on_close)
    root.mainloop()


# ------------------------------------------------------------------ config ----

CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")


def load_config():
    try:
        with open(CONFIG_PATH, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def save_config(cfg):
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(cfg, f, ensure_ascii=False, indent=2)
    except OSError:
        pass


# --------------------------------------------------------------- selftest ----


def selftest():
    def w(t, x0, x1, top):
        return {"text": t, "x0": x0, "x1": x1, "top": top, "bottom": top + 8}

    words = [w("Jan", 10, 40, 100), w("Kowalski", 100, 160, 101),
             w("12,50", 300, 330, 100),
             w("Ala", 10, 35, 130), w("Nowak", 100, 150, 130), w("7", 300, 310, 130)]
    rows = cluster_rows(words, ytol=3)
    assert len(rows) == 2, rows
    b = column_bounds(rows, 400, min_gap=6)
    assert len(b) == 2 and 40 < b[0] < 100 and 160 < b[1] < 300, b

    class P:
        width = 400

        def extract_words(self):
            return words

    assert rows_from_words(P(), 3, 6) == [["Jan", "Kowalski", "12,50"],
                                          ["Ala", "Nowak", "7"]]
    assert cast("1 234,56", True) == 1234.56
    assert cast("A1", True) == "A1"
    assert cast("0,50", True) == 0.5 and cast("-12.5", True) == -12.5
    for s in ("61109010140000071219812874", "00123", "007"):  # IBAN, kody
        assert cast(s, True) == s, cast(s, True)
    pr, rm = drop_repeats([[["H"], ["a"]], [["H"], ["b"]], [["H"], ["c"]]])
    assert pr == [[["a"]], [["b"]], [["c"]]] and rm == [["H"]] * 3, pr
    # powtarzajacy sie wiersz danych w srodku strony zostaje
    d = ["Dodatek", "200,00"]
    e = ["2", "Dodatek", "200,00"]  # ostatni wiersz danych pelnych stron
    pages = [[["Lista"], [n], d, [n + "2"], e, ["Str."]] for n in "XYZ"]
    pages.append([["Lista"], ["V"], d, ["Str."]])
    pr, rm = drop_repeats(pages)
    assert pr == [p[1:-1] for p in pages] and rm == [["Lista"], ["Str."]] * 4, pr
    print("selftest OK")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
    else:
        gui()
